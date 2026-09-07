"""Structural verification for the Review 1 deck.

LibreOffice is not installed here, so the deck cannot be rendered to images.
These checks are structural, not visual - they prove the file is well-formed and
that its animation targets resolve, not that it looks right. Opening it in
PowerPoint is still required.

    python presentation/verify_deck.py
"""
import os
import re
import sys
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content as C

HERE = os.path.dirname(os.path.abspath(__file__))
P = "http://schemas.openxmlformats.org/presentationml/2006/main"

ANIMATED = os.path.join(HERE, "CALIPER_Review1.pptx")
STATIC = os.path.join(HERE, "CALIPER_Review1_static.pptx")

fails, warns = [], []


def check(ok, msg):
    print(("  PASS  " if ok else "  FAIL  ") + msg)
    if not ok:
        fails.append(msg)


def warn(msg):
    print("  WARN  " + msg)
    warns.append(msg)


def slide_text(slide):
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
    return "\n".join(out)


# ------------------------------------------------------------------ 1. round-trip
print("\n1. round-trip and animation targets")
prs = Presentation(ANIMATED)
n_slides = len(prs.slides._sldIdLst)
check(n_slides >= 22, "slide count is %d (>= 22)" % n_slides)

bad_targets, animated_slides, total_steps = [], 0, 0
for i, slide in enumerate(prs.slides, 1):
    ids = {sh.shape_id for sh in slide.shapes}
    timings = slide._element.findall("{%s}timing" % P)
    if not timings:
        continue
    animated_slides += 1
    spids = {int(e.get("spid"))
             for t in timings for e in t.iter("{%s}spTgt" % P)}
    total_steps += len([1 for t in timings
                        for e in t.iter("{%s}cTn" % P)
                        if e.get("nodeType") == "clickEffect"])
    missing = spids - ids
    if missing:
        bad_targets.append((i, sorted(missing)))

check(not bad_targets,
      "every animation spid resolves to a shape on its own slide"
      + ("" if not bad_targets else " - broken: %r" % bad_targets))
print("        %d animated slides, %d click steps" % (animated_slides, total_steps))


# --------------------------------------------------------------- 2. xml position
print("\n2. timing XML placement")
order_ok, count_ok = True, True
with zipfile.ZipFile(ANIMATED) as z:
    names = [n for n in z.namelist()
             if re.match(r"ppt/slides/slide\d+\.xml$", n)]
    for n in sorted(names):
        root = etree.fromstring(z.read(n))
        kids = [etree.QName(k).localname for k in root
                if isinstance(k.tag, str)]
        if kids.count("timing") > 1:
            count_ok = False
        if "timing" in kids and kids[-1] != "timing":
            order_ok = False
        if "transition" in kids and "timing" in kids:
            if kids.index("transition") > kids.index("timing"):
                order_ok = False
check(count_ok, "at most one <p:timing> per slide")
check(order_ok, "<p:timing> is last, <p:transition> precedes it")


# ------------------------------------------------------------------ 3. zip health
print("\n3. package integrity")
for path, label in [(ANIMATED, "animated"), (STATIC, "static")]:
    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
        media = [n for n in z.namelist() if n.startswith("ppt/media/")]
        missing = []
        for rels in [n for n in z.namelist() if n.endswith(".rels")]:
            root = etree.fromstring(z.read(rels))
            base = os.path.dirname(os.path.dirname(rels))
            for r in root:
                if r.get("TargetMode") == "External":
                    continue
                tgt = r.get("Target")
                if tgt.startswith("../"):
                    resolved = os.path.normpath(
                        os.path.join(base, tgt)).replace("\\", "/")
                else:
                    resolved = os.path.normpath(
                        os.path.join(os.path.dirname(
                            os.path.dirname(rels)), tgt)).replace("\\", "/")
                if resolved not in z.namelist():
                    resolved2 = os.path.normpath(os.path.join(
                        os.path.dirname(rels).replace("/_rels", ""), tgt)
                    ).replace("\\", "/")
                    if resolved2 not in z.namelist():
                        missing.append((rels, tgt))
        check(bad is None, "%s: zip CRC clean" % label)
        check(not missing,
              "%s: every relationship target exists" % label
              + ("" if not missing else " - missing %r" % missing[:4]))
        print("        %s: %d embedded media files" % (label, len(media)))


# ---------------------------------------------------------------- 4. static deck
print("\n4. static fallback")
prs_s = Presentation(STATIC)
n_timing = sum(len(sl._element.findall("{%s}timing" % P)) for sl in prs_s.slides)
check(n_timing == 0, "static deck carries no timing trees")
check(len(prs_s.slides._sldIdLst) == n_slides,
      "static deck has the same %d slides" % n_slides)


# --------------------------------------------------------------- 5. text fidelity
print("\n5. canonical text fidelity")
all_text = "\n".join(slide_text(s) for s in prs.slides)
norm = lambda t: re.sub(r"\s+", " ", t).strip()
check(norm(C.PROBLEM_STATEMENT) in norm(all_text),
      "problem statement appears verbatim (docs/sem5-review-plan.md)")

src = os.path.join(os.path.dirname(HERE), "docs", "sem5-review-plan.md")
if os.path.exists(src):
    doc = open(src, encoding="utf-8").read()
    k = doc.find("Interpretability tools claim")
    raw = doc[k:doc.index(chr(10) + chr(10), k)] if k >= 0 else ""
    # strip blockquote markers and markdown emphasis, then normalise
    canon = norm(raw.replace("> ", "").replace("*", ""))
    tidy = lambda t: t.replace("’", "'").replace("—", "--")
    check(tidy(canon) == tidy(norm(C.PROBLEM_STATEMENT)),
          "content.py matches docs/sem5-review-plan.md word for word")
else:
    warn("sem5-review-plan.md not found; skipped source comparison")


# --------------------------------------------------------------- 6. results guard
print("\n6. results guard (Review 1 carries no results)")
BANNED = [
    (r"77\s*/\s*100", "the 77/100 gate result"),
    (r"\b77\s*%", "77 percent"),
    (r"\b23\s*%", "23 percent silent-failure rate"),
    (r"\b10\.[08]\s*%", "the 10.0/10.8 percent detection rate"),
    (r"9\.3\s*[eE×x]", "the p = 9.3e-9 result"),
    (r"\bWilcoxon\b", "the statistical test"),
    (r"\bp\s*[=<]\s*0?\.", "a p-value"),
    (r"\bour\s+(arXiv\s+)?preprint\b", "our preprint"),
    (r"\bPaper\s+A\b", "Paper A"),
    (r"\bStudy\s+3\s+(shows|found|result)", "a Study 3 result"),
]
hits = []
for pat, label in BANNED:
    for i, s in enumerate(prs.slides, 1):
        m = re.search(pat, slide_text(s), re.I)
        if m:
            hits.append("slide %d: %s (%r)" % (i, label, m.group(0)))
check(not hits, "no result numbers on any slide"
      + ("" if not hits else "\n          " + "\n          ".join(hits)))


# ------------------------------------------------------------- 7. overflow sweep
print("\n7. geometry sweep")
SW, SH = prs.slide_width, prs.slide_height
over = []
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        try:
            if sh.left is None or sh.top is None:
                continue
            r, b = sh.left + (sh.width or 0), sh.top + (sh.height or 0)
            if sh.left < -Emu(1) or sh.top < -Emu(1) or r > SW + Emu(9144) \
                    or b > SH + Emu(9144):
                over.append("slide %d: %s" % (i, sh.name))
        except Exception:
            pass
if over:
    for o in over[:12]:
        warn("shape past the slide edge - " + o)
else:
    print("  PASS  every shape sits inside the slide")

# placeholders still to fill
ph = [p for p in ("[PROJECT ID]", "[GUIDE NAME]", "[MEMBER B]",
                  "[MEMBER C]", "[MEMBER D]") if p in all_text]
if ph:
    warn("placeholders still to fill: " + ", ".join(ph))


# ------------------------------------------------------------------------ done
print("\n" + "=" * 62)
print("%d checks failed, %d warnings" % (len(fails), len(warns)))
if fails:
    for f in fails:
        print("  FAILED: " + f.splitlines()[0])
print("=" * 62)
sys.exit(1 if fails else 0)
